import json
import os

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt

SRC = os.environ.get("BI_SRC_DOCX", r"C:\Users\ADMIN\Downloads\BI&DSS\output_5.2\BI và DSS - hoàn thiện 5.2.docx")
OUT = os.environ.get("BI_OUT_DOCX", r"C:\Users\ADMIN\Downloads\bi_work\bi-dss-platform\output_6.2-6.3\BI và DSS - hoàn thiện 6.2-6.3.docx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)

R = json.load(open("out/raw_profile.json", encoding="utf8"))
M = json.load(open("out/mart_profile.json", encoding="utf8"))
X = json.load(open("out/export_manifest.json", encoding="utf8"))
C = json.load(open("out/contract_report.json", encoding="utf8"))
D = docx.Document(SRC)
SECT = D.element.body.find(qn("w:sectPr"))
W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

FM = M["flight_metrics"]
ST = {s["flight_status_code"]: s for s in M["status"]}
YC = {str(y["year"]): y for y in M["year_counts"]}
DT = M["delay_totals"]
DR = {r["delay_reason_code"]: r for r in M["delay_reasons"]}
RPF = {r["n_reasons"]: r["flights"] for r in M["reasons_per_flight"]}
TOTAL = FM["flights"]


def vn(x, d=0):
    s = f"{x:,.{d}f}"
    return s.replace(",", "§").replace(".", ",").replace("§", ".")


def pctf(x, d=2):
    return vn(x, d) + "%"


def para(text="", bold=False, italic=False, align=None, indent=True, size=None, font=None,
         after=152400, before=152400, keep=False, line=None):
    p = D.add_paragraph()
    SECT.addprevious(p._p)
    pf = p.paragraph_format
    pf.first_line_indent = Emu(457200 if indent else 0)
    pf.left_indent = Emu(0)
    pf.space_before = Emu(before)
    pf.space_after = Emu(after)
    if line:
        pf.line_spacing = line
    if align is not None:
        p.alignment = align
    if keep:
        pf.keep_with_next = True
    r = p.add_run(text)
    r.bold = bold or None
    r.italic = italic or None
    if size:
        r.font.size = Pt(size)
    if font:
        r.font.name = font
        r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    return p


def h(text):
    return para(text, bold=True, indent=False, keep=True, before=152400, after=0)


def sub(text):
    return para(text, bold=True, indent=False, before=60000, after=0, keep=True)


def cap(text):
    return para(text, italic=True, align=AL.CENTER, indent=False, after=60000, before=60000, keep=True)


def spacer():
    return para("", indent=False, before=0, after=0)


def bullet(text):
    p = para(text, indent=False, before=0, after=40000)
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    return p


def page_break():
    p = D.add_paragraph()
    SECT.addprevious(p._p)
    p.add_run().add_break(WD_BREAK.PAGE)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:fill"), fill)
    tcPr.append(s)


def table(rows, widths, size=10, align=None):
    t = D.add_table(rows=len(rows), cols=len(widths))
    SECT.addprevious(t._tbl)
    tblPr = t._tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblLayout", "w:tblBorders", "w:jc", "w:tblStyle"):
        for e in tblPr.findall(qn(tag)):
            tblPr.remove(e)
    borders = "".join(
        f'<w:{s} w:val="single" w:sz="8" w:space="0" w:color="000000"/>'
        for s in ("top", "left", "bottom", "right", "insideH", "insideV")
    )
    for xml in (
        f'<w:tblW {W_NS} w:w="{sum(widths)}" w:type="dxa"/>',
        f'<w:tblInd {W_NS} w:w="-480" w:type="dxa"/>',
        f"<w:tblBorders {W_NS}>{borders}</w:tblBorders>",
        f'<w:tblLayout {W_NS} w:type="fixed"/>',
    ):
        tblPr.append(docx.oxml.parse_xml(xml))
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    mar = (
        f'<w:tcMar {W_NS}><w:top w:w="60" w:type="dxa"/><w:left w:w="80" w:type="dxa"/>'
        '<w:bottom w:w="60" w:type="dxa"/><w:right w:w="80" w:type="dxa"/></w:tcMar>'
    )
    for i, row in enumerate(rows):
        trPr = t.rows[i]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            trPr.append(OxmlElement("w:tblHeader"))
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Emu(widths[j] * 635)
            c._tc.get_or_add_tcPr().append(docx.oxml.parse_xml(mar))
            p = c.paragraphs[0]
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            pf.first_line_indent = Emu(0)
            a = align[j] if align else "l"
            p.alignment = {"l": AL.LEFT, "c": AL.CENTER, "r": AL.RIGHT}[a]
            r = p.add_run(str(val))
            r.font.size = Pt(size)
            if i == 0:
                p.alignment = AL.CENTER
                r.bold = True
                shade(c, "D9D9D9")
    spacer()
    return t


def picture(path, width_cm=15.9):
    p = D.add_paragraph()
    SECT.addprevious(p._p)
    p.alignment = AL.CENTER
    pf = p.paragraph_format
    pf.first_line_indent = Emu(0)
    pf.keep_with_next = True
    pf.space_before = Pt(6)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    p.add_run().add_picture(path, width=Cm(width_cm))


def code(lines):
    for i, line in enumerate(lines):
        p = para(line, indent=False, size=8.5, font="Consolas", before=0, after=0,
                 align=AL.LEFT, line=1.0, keep=i < len(lines) - 1)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.keep_together = True
    spacer()


def y23(y):
    return " (01–08)" if y == "2023" else ""


page_break()

# ===================================================================== 6.2
h("6.2. Fact table(s)")
para(
    "Từ tầng staging đã chuẩn hóa, nhóm lọc dữ liệu về giai đoạn 2022–2023 và xây dựng hai bảng fact "
    "theo hai grain khác nhau: FactFlight (một dòng cho mỗi chuyến bay) và FactDelayReason (một dòng cho "
    "mỗi chuyến × nguyên nhân chậm có số phút dương). Việc lọc trước khi xuất CSV giúp Power BI chỉ nạp "
    f"{vn(TOTAL)} dòng FactFlight và {vn(DT['rows'])} dòng FactDelayReason thay vì toàn bộ 3.000.000 dòng, "
    "giảm dung lượng và thời gian refresh. Hai bảng được hiện thực bằng dbt-duckdb với contract enforced, "
    "được kiểm thử tự động và triển khai trên ai1-k3s."
)
picture("docs/figures/d2_data_flow.png")
cap("Hình 6.1. Luồng dữ liệu end-to-end từ CSV nguồn đến Power BI")

h("6.2.1. FactFlight")
para(
    "FactFlight là bảng fact chính ở grain một dòng cho mỗi chuyến bay. Khóa nghiệp vụ gồm năm trường "
    "(FL_DATE, AIRLINE_CODE, FL_NUMBER, ORIGIN, DEST); kiểm tra trên 1.151.344 chuyến cho kết quả 0 nhóm "
    "trùng và 0 dòng trùng, kể cả khi bổ sung CRS_DEP_TIME. FlightKey là khóa thay thế sinh bằng ROW_NUMBER "
    "theo đúng thứ tự khóa nghiệp vụ, bảo đảm tái lập qua các lần chạy."
)
para("Quy trình tạo bảng gồm năm bước:", before=100000)
bullet("Lọc FL_DATE từ 01/01/2022 đến 31/08/2023 (dữ liệu nguồn kết thúc tháng 8/2023).")
bullet("Phân loại trạng thái độc quyền theo thứ tự ưu tiên: Cancelled → Diverted → Delayed (ARR_DELAY ≥ 15) → On-time.")
bullet("Sinh DateKey dạng YYYYMMDD; map CarrierKey, OriginAirportKey, DestAirportKey, RouteKey và FlightStatusKey.")
bullet("Tính các measure/flag; chuẩn hóa giờ 2400 thành 0000; giữ NULL cho các measure không áp dụng thay vì điền 0.")
bullet("Xuất CSV UTF-8 cùng manifest checksum để nạp Power BI.")
cap("Bảng: Cấu trúc 28 thuộc tính của FactFlight")
rows = [["STT", "Thuộc tính", "Kiểu", "Mô tả", "Null"]]
flight_cols = [
    ("FlightKey", "BIGINT", "Khóa thay thế của chuyến bay."),
    ("DateKey", "INTEGER", "FK DimDate, dạng YYYYMMDD."),
    ("CarrierKey", "INTEGER", "FK DimCarrier."),
    ("OriginAirportKey", "INTEGER", "FK DimAirport cho sân bay đi."),
    ("DestAirportKey", "INTEGER", "FK DimAirport cho sân bay đến."),
    ("RouteKey", "INTEGER", "FK DimRoute."),
    ("FlightStatusKey", "INTEGER", "FK DimFlightStatus."),
    ("FlightNumber", "INTEGER", "Số hiệu chuyến (degenerate dimension)."),
    ("ScheduledDepTime", "INTEGER", "Giờ khởi hành theo lịch dạng hhmm; 2400 → 0000."),
    ("ScheduledDepHour", "INTEGER", "Giờ khởi hành theo lịch 0–23."),
    ("ScheduledDepMinute", "INTEGER", "Phút khởi hành theo lịch 0–59."),
    ("CancellationCode", "VARCHAR", "Mã lý do hủy A/B/C/D; NULL nếu không hủy."),
    ("DepDelay", "DOUBLE", "Chênh lệch giờ khởi hành thực tế so với lịch (phút, có thể âm)."),
    ("ArrDelay", "DOUBLE", "Chênh lệch giờ hạ cánh thực tế so với lịch (phút, có thể âm)."),
    ("DepDelayMinutes", "DOUBLE", "Số phút khởi hành muộn (≥ 0); NULL nếu hủy/thiếu."),
    ("ArrDelayMinutes", "DOUBLE", "Số phút hạ cánh muộn (≥ 0); NULL nếu hủy/chuyển hướng/thiếu."),
    ("TaxiOutMinutes", "DOUBLE", "Thời gian taxi ra (phút)."),
    ("TaxiInMinutes", "DOUBLE", "Thời gian taxi vào (phút)."),
    ("AirTimeMinutes", "DOUBLE", "Thời gian bay (phút)."),
    ("ElapsedTimeMinutes", "DOUBLE", "Thời gian hành trình thực tế (phút)."),
    ("CRSElapsedTimeMinutes", "DOUBLE", "Thời gian hành trình theo lịch (phút)."),
    ("DistanceMiles", "DOUBLE", "Khoảng cách bay (dặm)."),
    ("IsDepDelayed15", "INTEGER", "1 nếu DEP_DELAY ≥ 15 và chuyến không bị hủy."),
    ("IsArrDelayed15", "INTEGER", "1 nếu ARR_DELAY ≥ 15 trên chuyến hoàn thành."),
    ("IsCancelled", "INTEGER", "1 nếu chuyến bị hủy."),
    ("IsDiverted", "INTEGER", "1 nếu chuyến bị chuyển hướng."),
    ("IsCompleted", "INTEGER", "1 nếu không hủy và không chuyển hướng."),
    ("HasDelayReason", "INTEGER", "1 nếu có ít nhất một cột nguyên nhân chậm > 0."),
]
nullable = {"CancellationCode", "DepDelay", "ArrDelay", "DepDelayMinutes", "ArrDelayMinutes",
            "TaxiOutMinutes", "TaxiInMinutes", "AirTimeMinutes", "ElapsedTimeMinutes",
            "CRSElapsedTimeMinutes"}
for i, (name, typ, desc) in enumerate(flight_cols, 1):
    rows.append([i, name, typ, desc, "Có" if name in nullable else "Không"])
table(rows, [650, 2050, 1050, 5150, 1285], size=8, align=["c", "l", "c", "l", "c"])
picture("docs/figures/d2_star_flight.png")
cap("Hình 6.2. Star schema FactFlight — khóa chính, khóa ngoại và quan hệ")
cap("Bảng: Công thức các measure và cờ nghiệp vụ chính của FactFlight")
rows = [["Measure / cờ", "Công thức hoặc quy tắc", "Kiểu", "Ghi chú"]]
rows += [
    ["ArrDelayMinutes", "GREATEST(ARR_DELAY, 0) khi chuyến hoàn thành và ARR_DELAY không NULL", "DOUBLE",
     "Không áp dụng cho hủy/chuyển hướng."],
    ["DepDelayMinutes", "GREATEST(DEP_DELAY, 0) khi chuyến không hủy và DEP_DELAY không NULL", "DOUBLE",
     "Hủy chuyến → NULL."],
    ["IsArrDelayed15", "1 nếu IsCompleted = 1 và ARR_DELAY ≥ 15; ngược lại 0", "INTEGER",
     "Mẫu số phân tích là completed."],
    ["IsCancelled / IsDiverted", "Lấy trực tiếp từ CANCELLED / DIVERTED", "INTEGER",
     "Hai cờ không đồng thời bằng 1."],
    ["IsCompleted", "1 nếu CANCELLED = 0 và DIVERTED = 0; ngược lại 0", "INTEGER",
     "Dùng để tránh sai mẫu số."],
    ["FlightStatusKey", "Cancelled > Diverted > Delayed (ARR_DELAY ≥ 15) > On-time", "INTEGER",
     "Trạng thái độc quyền, map DimFlightStatus."],
    ["ScheduledDepTime", "hhmm; nếu bằng 2400 thì đổi thành 0000", "INTEGER",
     "Chuẩn hóa DQ08; tách thêm giờ/phút."],
]
table(rows, [2200, 5200, 900, 1885], size=8.5, align=["l", "l", "c", "l"])

cap("Bảng: Quy mô và trạng thái chuyến bay theo năm")
rows = [["Năm", "Số chuyến", "Đúng giờ", "Chậm ≥ 15", "Hủy", "Chuyển hướng", "Tỷ lệ chậm"]]
for y in ("2022", "2023"):
    v = YC[y]
    ontime = v["flights"] - v["cancelled"] - v["diverted"] - v["delayed15"]
    rows.append([y + y23(y), vn(v["flights"]), vn(ontime), vn(v["delayed15"]),
                 vn(v["cancelled"]), vn(v["diverted"]), pctf(v["delayed15"] / v["flights"] * 100)])
rows.append(["Tổng", vn(TOTAL), vn(ST["On-time"]["flights"]), vn(FM["delayed15"]),
             vn(FM["cancelled"]), vn(FM["diverted"]), pctf(FM["delayed15"] / TOTAL * 100)])
table(rows, [1400, 1500, 1400, 1500, 1100, 1400, 1885], size=9,
      align=["c", "r", "r", "r", "r", "r", "r"])

para(
    f"Trong {vn(TOTAL)} chuyến, có {vn(FM['delayed15'])} chuyến chậm ≥ 15 phút "
    f"({pctf(FM['delayed15'] / TOTAL * 100)}), {vn(FM['cancelled'])} chuyến bị hủy "
    f"({pctf(FM['cancelled'] / TOTAL * 100)}), {vn(FM['diverted'])} chuyến chuyển hướng "
    f"({pctf(FM['diverted'] / TOTAL * 100)}) và {vn(ST['On-time']['flights'])} chuyến đúng giờ "
    f"({pctf(ST['On-time']['flights'] / TOTAL * 100)}). ARR_DELAY trung bình là "
    f"{vn(FM['avg_arr_delay'], 2)} phút; median {vn(FM['median_arr_delay'])} phút; trung bình của nhóm chậm "
    f"là {vn(FM['avg_delayed_arr_delay'], 1)} phút. Có 2 chuyến hoàn thành nhưng thiếu ARR_DELAY; hai chuyến "
    "này được giữ ở trạng thái On-time với ArrDelayMinutes = NULL và được ghi nhận trong Data Quality Log.",
    before=100000,
)
para("Mã nguồn tiêu biểu (dbt, trích từ int_flight_enriched.sql):", bold=True, before=100000, keep=True)
code([
    "CASE WHEN cancelled = 1 THEN 'Cancelled'",
    "     WHEN diverted = 1 THEN 'Diverted'",
    "     WHEN arr_delay >= 15 THEN 'Delayed' ELSE 'On-time' END AS flight_status_code,",
    "ROW_NUMBER() OVER (ORDER BY flight_date, carrier_code, flight_number, origin_code, dest_code) AS flight_key,",
    "CAST(STRFTIME(flight_date, '%Y%m%d') AS INTEGER) AS date_key,",
    "CASE WHEN cancelled = 0 AND diverted = 0 AND arr_delay >= 15 THEN 1 ELSE 0 END AS is_arr_delayed15,",
    "CASE WHEN cancelled = 0 AND arr_delay IS NOT NULL THEN GREATEST(arr_delay, 0) END AS arr_delay_minutes",
])

h("6.2.2. FactDelayReason (tách riêng)")
para(
    "FactDelayReason được tách từ năm cột nguyên nhân của nguồn (DELAY_DUE_CARRIER, DELAY_DUE_WEATHER, "
    "DELAY_DUE_NAS, DELAY_DUE_SECURITY, DELAY_DUE_LATE_AIRCRAFT). Bảng có grain một dòng cho mỗi chuyến × "
    "nguyên nhân, chỉ giữ các dòng có DelayMinutes > 0. Kết quả có "
    f"{vn(DT['rows'])} dòng từ {vn(FM['delayed15'])} chuyến chậm, tổng "
    f"{vn(DT['minutes'])} phút, trung bình {vn(DT['minutes'] / FM['delayed15'], 2)} phút/chuyến chậm "
    f"({vn(DT['rows'] / FM['delayed15'], 2)} nguyên nhân/chuyến)."
)
para("Thiết kế giữ các khóa dùng chung DateKey, CarrierKey và OriginAirportKey; thêm FlightKey để drill-through "
     "về FactFlight và DelayReasonKey để phân tích nguyên nhân. Destination không đưa vào fact này theo yêu cầu "
     "tối giản khóa; khi cần phân tích theo điểm đến, dùng FlightKey hoặc bổ sung DestAirportKey trong phiên bản sau.",
     before=100000)
picture("docs/figures/d2_star_delay.png")
cap("Hình 6.3. Star schema FactDelayReason — khóa chính, khóa ngoại và quan hệ")
cap("Bảng: Cấu trúc FactDelayReason")
rows = [["STT", "Thuộc tính", "Kiểu", "Mô tả", "Null"]]
reason_cols = [
    ("DelayFactKey", "BIGINT", "Khóa thay thế của dòng fact."),
    ("FlightKey", "BIGINT", "FK FactFlight; drill-through."),
    ("DateKey", "INTEGER", "FK DimDate."),
    ("CarrierKey", "INTEGER", "FK DimCarrier."),
    ("OriginAirportKey", "INTEGER", "FK DimAirport cho sân bay đi."),
    ("DelayReasonKey", "INTEGER", "FK DimDelayReason (5 nhóm)."),
    ("DelayMinutes", "DOUBLE", "Số phút chậm do nguyên nhân tương ứng (> 0)."),
]
for i, (name, typ, desc) in enumerate(reason_cols, 1):
    rows.append([i, name, typ, desc, "Không"])
table(rows, [650, 2250, 1050, 4950, 1285], size=8.5, align=["c", "l", "c", "l", "c"])

cap("Bảng: Phân bố nguyên nhân chậm chuyến 2022–2023")
rows = [["Mã", "Tên nhóm", "Số dòng", "Tổng phút", "Phút TB/dòng", "Tỷ trọng phút"]]
for rcode in ("CARRIER", "WEATHER", "NAS", "SECURITY", "LATE_AIRCRAFT"):
    v = DR[rcode]
    rows.append([rcode, v["delay_reason_name"], vn(v["rows"]), vn(v["minutes"]),
                 vn(v["avg_minutes"], 2), pctf(v["share_pct"])])
rows.append(["Tổng", "5 nhóm", vn(DT["rows"]), vn(DT["minutes"]),
             vn(DT["avg_minutes"], 2), "100,00%"])
table(rows, [1800, 2500, 1300, 1600, 1400, 1585], size=8.5,
      align=["l", "l", "r", "r", "r", "r"])
para(
    "Hai nhóm LATE_AIRCRAFT và CARRIER chiếm 76,89% tổng số phút chậm; NAS chiếm 17,46%; WEATHER 5,44%; "
    "SECURITY chỉ 0,21%. Số dòng của từng nhóm không cộng thành số chuyến vì một chuyến có thể có nhiều nguyên nhân.",
    before=100000,
)

cap("Bảng: Số nguyên nhân trên mỗi chuyến chậm")
rows = [["Số nguyên nhân/chuyến", "Số chuyến", "Tỷ trọng"]]
for n in (1, 2, 3, 4, 5):
    v = RPF.get(n, 0)
    rows.append([n, vn(v), pctf(v / FM["delayed15"] * 100)])
rows.append(["Tổng", vn(FM["delayed15"]), "100,00%"])
table(rows, [4000, 3000, 3185], size=9, align=["c", "r", "r"])

para(
    "Đối soát nghiệp vụ là điểm mạnh của thiết kế này: với cả 245.127 chuyến có ARR_DELAY ≥ 15, tổng năm "
    "cột nguyên nhân bằng đúng ARR_DELAY (chênh lệch tối đa 0 phút). Không có nguyên nhân nào xuất hiện trên "
    "chuyến đúng giờ/sớm, chuyến hủy hoặc chuyển hướng; không có chuyến chậm ≥ 15 nào thiếu nguyên nhân. "
    "Độ trễ lớn nhất của một dòng nguyên nhân là 2.934 phút; đây là giá trị hợp lệ nghiệp vụ và được giữ lại.",
    before=100000,
)
para("Mã nguồn tiêu biểu (dbt, fact_delay_reason.sql):", bold=True, before=100000, keep=True)
code([
    "WITH unpivoted AS (",
    "  SELECT flight_key, date_key, carrier_key, origin_airport_key,",
    "         'CARRIER' AS delay_reason_code, delay_due_carrier AS delay_minutes FROM enriched",
    "  UNION ALL SELECT flight_key, date_key, carrier_key, origin_airport_key, 'WEATHER', delay_due_weather FROM enriched",
    "  UNION ALL SELECT flight_key, date_key, carrier_key, origin_airport_key, 'NAS', delay_due_nas FROM enriched",
    "  UNION ALL SELECT flight_key, date_key, carrier_key, origin_airport_key, 'SECURITY', delay_due_security FROM enriched",
    "  UNION ALL SELECT flight_key, date_key, carrier_key, origin_airport_key, 'LATE_AIRCRAFT', delay_due_late_aircraft FROM enriched",
    ")",
    "SELECT ROW_NUMBER() OVER (ORDER BY flight_key, delay_reason_code) AS delay_fact_key,",
    "       u.flight_key, u.date_key, u.carrier_key, u.origin_airport_key,",
    "       d.delay_reason_key, u.delay_minutes",
    "FROM unpivoted u JOIN dim_delay_reason d USING (delay_reason_code)",
    "WHERE u.delay_minutes > 0",
])

h("6.2.3. Ghi chú về việc tách fact table")
para(
    "Việc tách thành hai fact không phải để trang trí mô hình mà xuất phát từ sự khác biệt về grain và cách "
    "cộng dồn measure. Nếu gộp năm cột nguyên nhân vào FactFlight, bảng sẽ rộng, phần lớn là NULL (chỉ 21,29% "
    "chuyến chậm mới có nguyên nhân), và mọi measure cấp chuyến sẽ bị lặp lại khi unpivot. Ngược lại, nếu đưa "
    "các measure cấp chuyến vào FactDelayReason, một chuyến có ba nguyên nhân sẽ bị đếm ba lần. Hai grain phải "
    "nằm ở hai bảng."
)
cap("Bảng: So sánh hai bảng fact")
rows = [["Tiêu chí", "FactFlight", "FactDelayReason", "Hệ quả thiết kế"],
        ["Grain", "1 dòng / chuyến", "1 dòng / chuyến × nguyên nhân", "Không trộn grain"],
        ["Số dòng", vn(TOTAL), vn(DT["rows"]), "Fact phụ nhỏ hơn ~2,9 lần"],
        ["Measure chính", "ArrDelayMinutes, DepDelayMinutes, IsArrDelayed15, IsCancelled", "DelayMinutes theo 5 nhóm nguyên nhân", "Measure additive đúng cấp"],
        ["Dimension dùng chung", "Date, Carrier, Origin, Dest, Route, Status", "Date, Carrier, Origin, DelayReason + FlightKey", "Conformed dimensions"],
        ["Câu hỏi phân tích", "Tỷ lệ chậm/hủy, xu hướng, tuyến, hãng, sân bay", "Nguyên nhân nào chiếm nhiều phút nhất", "Hai fact phục vụ hai nhóm câu hỏi"],
        ["Power BI", "Fact chính, quét 1,15 triệu dòng", "Fact phụ, quét 397 nghìn dòng", "Giảm thời gian lọc nguyên nhân"]]
table(rows, [1900, 3100, 3200, 1985], size=8.5, align=["l", "l", "l", "l"])
para(
    "Chi tiết lợi ích: (i) NULL density thấp hơn nhiều so với một bảng rộng; (ii) phân tích nguyên nhân chỉ "
    "quét 397.131 dòng thay vì 1,15 triệu; (iii) mỗi fact có contract và test riêng; (iv) tránh double-counting "
    "khi tính tỷ lệ chậm theo hãng/tuyến; (v) dễ mở rộng nếu BTS bổ sung nhóm nguyên nhân mới. Chi phí là mô hình "
    "Power BI có thêm một fact và cần drill-through qua FlightKey khi muốn kết hợp nguyên nhân với điểm đến; "
    "đây là chi phí hợp lý so với sai số do trộn grain.",
    before=100000,
)
para(
    "Phương án đã cân nhắc nhưng không chọn: (a) một bảng fact rộng với năm cột nguyên nhân — đơn giản nhưng "
    "nhiều NULL và dễ double-count khi unpivot tại DAX; (b) một bảng cầu nối nguyên nhân — thêm độ phức tạp "
    "không cần thiết cho 5 nhóm cố định; (c) giữ cả measure cấp chuyến trong FactDelayReason — vi phạm grain. "
    "Thiết kế hai fact + conformed dimensions là lựa chọn cân bằng nhất cho cả phân tích mô tả và dashboard.",
    before=100000,
)

# ===================================================================== 6.3
h("6.3. Dimension table(s)")
para(
    "Sáu bảng chiều được tạo sau khi lọc fact, bảo đảm mỗi giá trị khóa trong fact đều có dòng chiều tương ứng. "
    "Các bảng dùng khóa thay thế kiểu số nguyên, giữ khóa tự nhiên để truy vết, loại bỏ trùng lặp và chuẩn hóa "
    "định dạng. DimDate, DimCarrier và DimAirport là ba chiều dùng chung cho cả hai fact; DimRoute và "
    "DimFlightStatus phục vụ FactFlight; DimDelayReason phục vụ FactDelayReason."
)
cap("Bảng: Tổng hợp sáu bảng chiều")
rows = [["Bảng", "Grain", "Số dòng", "Khóa chính", "Khóa tự nhiên", "Nguồn / ghi chú"],
        ["DimDate", "1 dòng / ngày", "730", "DateKey", "FullDate", "Sinh 2022-01-01…2023-12-31"],
        ["DimCarrier", "1 dòng / hãng", "17", "CarrierKey", "CarrierCode", "AIRLINE_CODE; 1–1 với tên/DOT"],
        ["DimAirport", "1 dòng / sân bay", "370", "AirportKey", "AirportCode", "OurAirports + Census; 2 override"],
        ["DimRoute", "1 dòng / tuyến", "6.820", "RouteKey", "OriginCode + DestCode", "DistanceGroup; 80 tuyến có 2 distance"],
        ["DimDelayReason", "1 dòng / nhóm nguyên nhân", "5", "DelayReasonKey", "DelayReasonCode", "5 cột DELAY_DUE_*"],
        ["DimFlightStatus", "1 dòng / trạng thái", "4", "FlightStatusKey", "FlightStatusCode", "Trạng thái độc quyền"]]
table(rows, [1700, 2200, 900, 1500, 1900, 1985], size=8.5, align=["l", "l", "r", "l", "l", "l"])

h("6.3.1. DimDate")
para(
    "DimDate được sinh đủ 730 ngày từ 01/01/2022 đến 31/12/2023, gồm 365 ngày mỗi năm; mỗi năm có 105 ngày "
    "cuối tuần (tổng 210). Bảng vẫn phủ đủ năm 2023 dù fact chỉ có dữ liệu đến 31/08/2023, để dashboard có trục "
    "thời gian liên tục và tránh lỗi lọc tháng."
)
cap("Bảng: Cấu trúc DimDate")
rows = [["STT", "Thuộc tính", "Kiểu", "Mô tả"]]
date_cols = [
    ("DateKey", "INTEGER", "Khóa ngày dạng YYYYMMDD."),
    ("FullDate", "DATE", "Ngày đầy đủ."),
    ("Year", "INTEGER", "Năm 2022/2023."),
    ("Quarter", "INTEGER", "Quý 1–4."),
    ("QuarterName", "VARCHAR", "Q1…Q4."),
    ("Month", "INTEGER", "Tháng 1–12 (dùng để sort)."),
    ("MonthName", "VARCHAR", "Tháng 1…Tháng 12."),
    ("MonthYear", "VARCHAR", "Nhãn YYYY-MM."),
    ("DayOfMonth", "INTEGER", "Ngày trong tháng."),
    ("DayOfYear", "INTEGER", "Ngày trong năm."),
    ("WeekOfYear", "INTEGER", "Tuần ISO."),
    ("WeekdayNumber", "INTEGER", "ISO: 1 = Thứ Hai, 7 = Chủ Nhật."),
    ("WeekdayName", "VARCHAR", "Thứ Hai…Chủ Nhật."),
    ("IsWeekend", "BOOLEAN", "TRUE nếu Thứ Bảy/Chủ Nhật."),
    ("Season", "VARCHAR", "Xuân, Hạ, Thu, Đông theo tháng."),
]
for i, (name, typ, desc) in enumerate(date_cols, 1):
    rows.append([i, name, typ, desc])
table(rows, [650, 2100, 1050, 6385], size=8.5, align=["c", "l", "c", "l"])
para(
    f"DimDate có {vn(sum(x['days'] for x in M['date_months']))} dòng; năm 2022 và 2023 đều có 365 ngày, "
    "105 ngày cuối tuần mỗi năm. Season được gán theo tháng: tháng 3–5 Xuân, 6–8 Hạ, 9–11 Thu, 12–2 Đông.",
    before=100000,
)

h("6.3.2. DimCarrier")
para(
    "DimCarrier lấy từ bốn trường AIRLINE_CODE, AIRLINE, DOT_CODE và AIRLINE_DOT. Sau khi lọc 2022–2023 còn 17 "
    "hãng (toàn bộ dữ liệu có 18 hãng; một hãng không còn hoạt động trong giai đoạn này). Quan hệ giữa mã hãng "
    "và tên/DOT là 1–1; không có mã nào ứng với hai tên hoặc ngược lại."
)
cap("Bảng: Cấu trúc DimCarrier")
rows = [["STT", "Thuộc tính", "Kiểu", "Mô tả"]]
for i, (name, typ, desc) in enumerate([
    ("CarrierKey", "INTEGER", "Khóa thay thế."),
    ("CarrierCode", "VARCHAR", "Mã BTS (AIRLINE_CODE)."),
    ("CarrierName", "VARCHAR", "Tên hãng (AIRLINE)."),
    ("DotCode", "INTEGER", "Mã DOT."),
    ("AirlineDot", "VARCHAR", "Chuỗi AIRLINE_DOT gốc."),
], 1):
    rows.append([i, name, typ, desc])
table(rows, [650, 2100, 1050, 6385], size=8.5, align=["c", "l", "c", "l"])
cap("Bảng: Danh sách 17 hãng trong giai đoạn 2022–2023")
rows = [["Mã", "Tên hãng", "Số chuyến", "Chậm ≥ 15", "Tỷ lệ hủy"]]
for c in M["top_carriers"]:
    rows.append([c["carrier_code"], c["carrier_name"], vn(c["flights"]),
                 pctf(c["delayed15_pct"]), pctf(c["cancelled_pct"])])
table(rows, [1200, 4200, 1500, 1700, 1585], size=8, align=["c", "l", "r", "r", "r"])
para(
    "Southwest dẫn đầu về số chuyến (230.123; 24,07% chậm ≥ 15); Frontier và Allegiant có tỷ lệ chậm cao "
    "nhất (34,24% và 31,34%); JetBlue 32,84%. Các hãng khu vực có tỷ lệ hủy cao hơn mặt bằng chung (Republic 4,16%).",
    before=100000,
)

h("6.3.3. DimAirport")
para(
    "DimAirport có 370 sân bay xuất hiện trong fact 2022–2023 (toàn bộ dữ liệu có 380 mã). Từ điển nguồn chỉ có "
    "mã, tên thành phố và bang dạng chuỗi; nhóm làm giàu thêm AirportName, State, Region, Country và tọa độ từ "
    "OurAirports (public domain) kết hợp Census region. Trong 370 sân bay, 368 khớp trực tiếp OurAirports và 2 "
    "được override thủ công: PBI (OurAirports 2026 đã đổi mã thành DJT) và ISN (không còn trong OurAirports). "
    "Các lãnh thổ PR, VI, GU, AS, MP được map về đúng mã và vùng; không còn sân bay nào thiếu State/Region."
)
cap("Bảng: Cấu trúc DimAirport")
rows = [["STT", "Thuộc tính", "Kiểu", "Mô tả"]]
for i, (name, typ, desc) in enumerate([
    ("AirportKey", "INTEGER", "Khóa thay thế."),
    ("AirportCode", "VARCHAR", "Mã IATA 3 ký tự."),
    ("AirportName", "VARCHAR", "Tên sân bay."),
    ("City", "VARCHAR", "Thành phố."),
    ("StateCode", "VARCHAR", "Mã bang/lãnh thổ 2 ký tự."),
    ("StateName", "VARCHAR", "Tên bang/lãnh thổ."),
    ("Region", "VARCHAR", "Northeast, Midwest, South, West."),
    ("RegionName", "VARCHAR", "Tên vùng tiếng Việt."),
    ("CountryCode", "VARCHAR", "Mã quốc gia ISO-2."),
    ("Latitude", "DOUBLE", "Vĩ độ."),
    ("Longitude", "DOUBLE", "Kinh độ."),
    ("DataSource", "VARCHAR", "ourairports / manual / fallback / missing."),
], 1):
    rows.append([i, name, typ, desc])
table(rows, [650, 2100, 1050, 6385], size=8.5, align=["c", "l", "c", "l"])
cap("Bảng: Phân bố sân bay và chuyến bay theo vùng")
rows = [["Vùng", "Số sân bay", "Số chuyến (origin)", "Tỷ lệ chậm ≥ 15", "Tỷ lệ hủy"]]
reg_flights = {r["region"]: r for r in M["regions"]}
for r in M["airports_by_region"]:
    f = reg_flights[r["region"]]
    rows.append([r["region"], vn(r["airports"]), vn(f["flights"]),
                 pctf(f["delayed15_pct"]), pctf(f["cancelled_pct"])])
table(rows, [1800, 1900, 2100, 2200, 2185], size=9, align=["l", "r", "r", "r", "r"])
para(
    "South dẫn đầu với 124 sân bay và 466.979 chuyến; West 117 sân bay/357.987 chuyến; Midwest 90/176.871; "
    "Northeast 39/149.507. Tỷ lệ chậm cao nhất ở Northeast (23,14%) và hủy cao nhất cũng ở Northeast (3,49%). "
    "Các bang nhiều sân bay nhất: TX 26, CA 22, AK 19, FL 19, MI 15, NY 15.",
    before=100000,
)
cap("Bảng: 10 sân bay đi lớn nhất theo số chuyến")
rows = [["Mã", "Sân bay", "Số chuyến", "Chậm ≥ 15", "Tỷ lệ hủy"]]
for a in M["top_origins"][:10]:
    rows.append([a["airport_code"], a["airport_name"], vn(a["flights"]),
                 pctf(a["delayed15_pct"]), pctf(a["cancelled_pct"])])
table(rows, [900, 4500, 1400, 1700, 1685], size=8, align=["c", "l", "r", "r", "r"])
para(
    "ATL dẫn đầu với 55.262 chuyến đi; DFW 47.348; DEN 47.012; ORD 43.907. DEN có tỷ lệ chậm 27,48%, "
    "LAS 28,20% — cao hơn mặt bằng chung 21,29%; LGA có tỷ lệ hủy 4,43%.",
    before=100000,
)

h("6.3.4. DimRoute")
para(
    "DimRoute gộp theo cặp ORIGIN–DEST, sinh RouteKey, map hai FK sân bay và chia DistanceGroup. Có 6.820 tuyến; "
    "80 tuyến có hai giá trị DISTANCE khác nhau (tối đa 2 giá trị), do chênh lệch làm tròn/đường bay; "
    "distance_miles dùng trung bình của tuyến. Nhóm khoảng cách được dùng để so sánh tỷ lệ chậm giữa chặng ngắn, "
    "trung bình và dài."
)
cap("Bảng: Cấu trúc DimRoute")
rows = [["STT", "Thuộc tính", "Kiểu", "Mô tả"]]
for i, (name, typ, desc) in enumerate([
    ("RouteKey", "INTEGER", "Khóa thay thế của tuyến."),
    ("OriginCode", "VARCHAR", "Mã sân bay đi."),
    ("DestCode", "VARCHAR", "Mã sân bay đến."),
    ("RouteName", "VARCHAR", "Nhãn ORIGIN-DEST."),
    ("OriginAirportKey", "INTEGER", "FK DimAirport."),
    ("DestAirportKey", "INTEGER", "FK DimAirport."),
    ("DistanceMiles", "DOUBLE", "Khoảng cách trung bình (dặm)."),
    ("DistanceGroup", "VARCHAR", "0–499, 500–999, 1000–1499, 1500–1999, 2000+."),
    ("DistanceGroupOrder", "INTEGER", "Thứ tự nhóm 1–5."),
], 1):
    rows.append([i, name, typ, desc])
table(rows, [650, 2100, 1050, 6385], size=8.5, align=["c", "l", "c", "l"])
cap("Bảng: Phân bố chuyến bay và tỷ lệ chậm theo nhóm khoảng cách")
rows = [["Nhóm khoảng cách", "Số chuyến", "Tỷ trọng", "ARR_DELAY TB", "Tỷ lệ chậm ≥ 15"]]
for g in M["distance_groups"]:
    rows.append([g["distance_group"], vn(g["flights"]), pctf(g["flights"] / TOTAL * 100),
                 vn(g["avg_arr_delay"], 2), pctf(g["delayed15_pct"])])
table(rows, [2400, 1600, 1600, 2000, 2585], size=9, align=["l", "r", "r", "r", "r"])
para(
    "Nhóm 1.000–1.499 dặm có tỷ lệ chậm cao nhất (24,99%) và ARR_DELAY trung bình 10,45 phút; nhóm dưới "
    "500 dặm thấp nhất (19,44%, 6,18 phút). Hai nhóm 0–499 và 500–999 chiếm 71,49% tổng số chuyến.",
    before=100000,
)

h("6.3.5. DimDelayReason")
para(
    "DimDelayReason chuẩn hóa năm cột nguyên nhân của BTS thành năm dòng có mã, tên tiếng Việt, mô tả và "
    "nguồn cột. Đây là chiều tĩnh, được version hóa trong dbt; mọi dòng FactDelayReason đều map được sang đúng "
    "một DelayReasonKey."
)
cap("Bảng: Năm nhóm nguyên nhân chậm chuyến")
rows = [["Key", "Mã", "Tên nhóm", "Mô tả", "Cột nguồn"]]
for r in M["samples"]["dim_delay_reason"]:
    rows.append([r["delay_reason_key"], r["delay_reason_code"], r["delay_reason_name"],
                 r["description"], r["source_column"]])
table(rows, [700, 1700, 2300, 3685, 1800], size=8.5, align=["c", "l", "l", "l", "l"])
para(
    "Tổng của năm nhóm bằng ARR_DELAY trên toàn bộ 245.127 chuyến chậm; vì vậy DelayMinutes có thể cộng trực "
    "tiếp trong FactDelayReason. Tên tiếng Việt giúp dashboard thân thiện, nhưng mã gốc vẫn được giữ để đối chiếu BTS.",
    before=100000,
)

h("6.3.6. DimFlightStatus")
para(
    "DimFlightStatus định nghĩa bốn trạng thái độc quyền. Thứ tự ưu tiên khi phân loại là Cancelled → Diverted → "
    "Delayed (ARR_DELAY ≥ 15) → On-time. Vì vậy một chuyến hủy không bao giờ bị đếm là chậm; một chuyến chuyển "
    "hướng cũng không vào nhóm chậm. On-time bao gồm cả chuyến đến sớm và chuyến chậm dưới 15 phút."
)
cap("Bảng: Bốn trạng thái chuyến bay và phân bố trên FactFlight")
rows = [["Key", "Mã", "Tên", "Định nghĩa", "Completed", "Disrupted", "Số chuyến (tỷ trọng)"]]
for s in M["samples"]["dim_flight_status"]:
    st = ST[s["flight_status_code"]]
    rows.append([s["flight_status_key"], s["flight_status_code"], s["flight_status_name"],
                 s["description"], "Có" if s["is_completed"] else "Không",
                 "Có" if s["is_disrupted"] else "Không",
                 f"{vn(st['flights'])} ({pctf(st['share_pct'])})"])
table(rows, [600, 1300, 1600, 3085, 1000, 1000, 1600], size=8.5,
      align=["c", "l", "l", "l", "c", "c", "r"])
para(
    f"On-time {vn(ST['On-time']['flights'])} ({pctf(ST['On-time']['share_pct'])}); Delayed "
    f"{vn(ST['Delayed']['flights'])} ({pctf(ST['Delayed']['share_pct'])}); Cancelled "
    f"{vn(ST['Cancelled']['flights'])} ({pctf(ST['Cancelled']['share_pct'])}); Diverted "
    f"{vn(ST['Diverted']['flights'])} ({pctf(ST['Diverted']['share_pct'])}). Hai chuyến hoàn thành nhưng thiếu "
    "ARR_DELAY được xếp On-time; khi tính tỷ lệ chậm nên dùng IsCompleted làm mẫu số để tránh sai lệch.",
    before=100000,
)

# ===================================================================== deployment note
h("Ghi chú triển khai trên ai1-k3s")
picture("docs/figures/d2_deployment.png")
cap("Hình 6.4. Kiến trúc triển khai BI & DSS trên ai1-k3s")
para(
    "Mô hình được triển khai theo kiến trúc GitOps: mã nguồn ở GitHub (Royal2005-coder/bi-dss-platform), Flux CD "
    "theo dõi repo và áp Kustomization vào namespace bi-dss; Job bi-dss-build-v2 chạy dbt build, export CSV và "
    "kiểm tra contract; CronJob bi-dss-refresh (đang suspend) dùng cho lịch làm mới. Dữ liệu thô 586 MB và "
    "warehouse DuckDB nằm trên PVC local-path 20 GiB, không đưa vào Git. Image bi-dss-toolbox:1.0.0 "
    "(dbt-core 1.12.6, dbt-duckdb 1.11.0, DuckDB 1.5.6) được build local và import vào containerd k8s.io."
)
para(
    "Kết quả chạy chính thức trên cluster: dbt build PASS 145/145 trong 6,34 giây; Job hoàn tất sau 23 giây; "
    f"{C['checks']} kiểm tra contract, {C['failures']} lỗi. Tám CSV được xuất vào /data/exports/powerbi; "
    "checksum SHA-256 của FactFlight, FactDelayReason và các dimension lớn khớp tuyệt đối với lần chạy host, "
    "xác nhận tính tái lập. Power BI nạp 8 CSV, tạo 12 quan hệ active và 1 quan hệ drill-through tùy chọn "
    "FactDelayReason → FactFlight qua FlightKey; FactFlight là fact chính, FactDelayReason là fact phụ.",
    before=100000,
)
cap("Bảng: Quan hệ Power BI xuất kèm (powerbi_relationships.csv)")
rows = [["From", "FromColumn", "To", "ToColumn", "Cardinality", "Active"]]
for r in X["relationships"]:
    rows.append([r["from_table"], r["from_column"], r["to_table"], r["to_column"],
                 "N:1", "Có" if r["active"] else "Không (tùy chọn)"])
table(rows, [1700, 1700, 1700, 1700, 1400, 1985], size=8, align=["l", "l", "l", "l", "c", "c"])
para(
    "Vận hành: khi dữ liệu nguồn thay đổi, thay file trong /data/raw qua loader pod rồi bật CronJob hoặc tạo Job "
    "mới; khi model thay đổi, tăng tên Job (bi-dss-build-v3) để tránh immutable Job; khi dependency thay đổi, "
    "rebuild image và import lại. Dữ liệu thô, warehouse, log và CSV không nằm trong Git; chỉ code, contract và "
    "manifest được version hóa.",
    before=100000,
)
para(
    "Kết luận: mô hình hai fact + sáu dimension đã hoàn chỉnh về grain, khóa, measure và quan hệ; dữ liệu "
    "2022–2023 đã được kiểm thử tự động 145 test và 185 contract check, đạt mức sẵn sàng để nạp Power BI. "
    "Các giới hạn đã ghi nhận gồm 2 chuyến thiếu ARR_DELAY, 80 tuyến có hai giá trị khoảng cách, việc không "
    "đưa DestAirportKey vào FactDelayReason và CronJob đang tạm dừng — tất cả đều có xử lý hoặc hướng mở rộng rõ ràng.",
    before=100000,
)

D.save(OUT)
print("saved", OUT)
