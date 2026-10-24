#!/usr/bin/env python3
"""Build the Vietnamese handoff from README, DAX and optional observed JSON evidence."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections.abc import Mapping
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REPORT = ROOT / "outputs" / "validation_report.json"
DEFAULT_OUTPUT = ROOT / "outputs" / "SkyOps_BanGiao_KyThuat_6.2-6.6.docx"
DAX_FILES = (
    ROOT / "powerbi" / "dax_measures" / "01_core_kpis.dax",
    ROOT / "powerbi" / "dax_measures" / "02_delay_metrics.dax",
)


def set_font(style, name: str, size: float) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def configure(document) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.2), Cm(2)
    section.left_margin, section.right_margin = Cm(2.7), Cm(2)
    normal = document.styles["Normal"]
    set_font(normal, "Times New Roman", 12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
        style = document.styles[name]
        set_font(style, "Times New Roman", size)
        style.font.color.rgb = RGBColor.from_string("18364A")
        style.paragraph_format.keep_with_next = True
    set_font(document.styles["Caption"], "Times New Roman", 10)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("SkyOps — Bàn giao kỹ thuật | ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = "SkyOps — Bàn giao kỹ thuật 6.2–6.6"
    document.core_properties.subject = "ETL Pandas, mô hình dữ liệu và DAX cho Phúc"
    document.core_properties.language = "vi-VN"


def add_inline(paragraph, text: str) -> None:
    for token in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not token:
            continue
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            paragraph.add_run(token)


def add_code(document, text: str) -> None:
    for line in text.splitlines():
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = ([8.0, 3.0, 3.0, 2.3] if headers[0] == "Kiểm tra"
              else [4.0, 12.3] if len(headers) == 2
              else [16.3 / len(headers)] * len(headers))
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        add_inline(cell.paragraphs[0], str(header))
        for run in cell.paragraphs[0].runs:
            run.bold = True
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "E7EDF2")
        cell._tc.get_or_add_tcPr().append(shade)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if index < len(cells):
                add_inline(cells[index].paragraphs[0], str(value))
    for row in table.rows:
        # Keep a single table row together, but allow long tables to span pages.
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for cell_index, cell in enumerate(row.cells):
            cell.width = Cm(widths[cell_index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def render_markdown(document, source: str) -> None:
    lines = source.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("```"):
            index += 1
            block = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            add_code(document, "\n".join(block))
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            parsed = [[cell.strip() for cell in row.strip("|").split("|")] for row in table_lines]
            rows = [row for row in parsed[1:] if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
            add_table(document, parsed[0], rows)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            depth, text = len(heading.group(1)), heading.group(2)
            paragraph = document.add_heading(level=0 if depth == 1 else min(depth - 1, 3))
            add_inline(paragraph, text)
            index += 1
            continue
        block = [line]
        index += 1
        while index < len(lines) and lines[index].strip() and not lines[index].lstrip().startswith(("#", "|", "```")):
            block.append(lines[index].strip())
            index += 1
        add_inline(document.add_paragraph(), " ".join(block))


def display(value) -> str:
    if value is None:
        return "null (theo báo cáo)"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value)


def metric_rows(value, prefix: str = ""):
    if isinstance(value, Mapping):
        if not value:
            yield [prefix, "{}"]
        for key, child in value.items():
            yield from metric_rows(child, f"{prefix}.{key}" if prefix else str(key))
    elif isinstance(value, list):
        if not value:
            yield [prefix, "[]"]
        for index, child in enumerate(value):
            yield from metric_rows(child, f"{prefix}[{index}]")
    else:
        yield [prefix, display(value)]


def first(record, names, default="Không có trong báo cáo"):
    for name in names:
        if name in record:
            return record[name]
    return default


def add_evidence(document, report_path: Path, report) -> None:
    document.add_page_break()
    document.add_heading("Phụ lục A. Bằng chứng kiểm tra được cung cấp", level=1)
    if report is None:
        document.add_paragraph("Chưa có validation_report.json để đính kèm. Không có số dòng/KPI hoặc kết luận PASS được suy đoán. DAX và Power BI chưa được xác nhận thực thi từ tài liệu này.")
        return
    document.add_paragraph(f"Nguồn: {report_path.name}; SHA-256: {hashlib.sha256(report_path.read_bytes()).hexdigest()}")
    document.add_paragraph("Các giá trị dưới đây được trích từ JSON, không được generator đo lại. Status của report chỉ có hiệu lực cho phạm vi kiểm tra mà report mô tả; không xác nhận Power BI.")
    summary = [["status", display(report.get("status", "Không có trong báo cáo"))]]
    for key in ("generated_at", "generated_at_utc", "timestamp", "source", "raw_source", "dbt", "contracts"):
        if key in report:
            summary.extend(metric_rows(report[key], key))
    checks = report.get("checks")
    if isinstance(checks, list):
        passed = sum(isinstance(item, Mapping) and item.get("passed") is True for item in checks)
        failed = sum(isinstance(item, Mapping) and item.get("passed") is False for item in checks)
        summary.extend([
            ["Số mục checks có trong JSON", str(len(checks))],
            ["Số kiểm tra PASS", str(passed)],
            ["Số kiểm tra FAIL", str(failed)],
            ["Số mục thiếu passed Boolean", str(len(checks) - passed - failed)],
        ])
    add_table(document, ["Thuộc tính", "Giá trị quan sát trong báo cáo"], summary)
    tables = report.get("tables")
    if isinstance(tables, Mapping):
        entries = [(str(name), item) for name, item in tables.items()]
    elif isinstance(tables, list):
        entries = [(str(first(item, ("name", "table", "file"), f"Bảng {index + 1}")), item) for index, item in enumerate(tables) if isinstance(item, Mapping)]
    else:
        entries = []
    if entries:
        document.add_heading("Số liệu tám bảng / file được báo cáo", level=2)
        rows, hashes, schemas = [], [], []
        for name, item in entries:
            if not isinstance(item, Mapping):
                rows.append([name, display(item), "—", "—"])
                continue
            columns = first(item, ("columns", "column_count", "n_columns"))
            column_count = len(columns) if isinstance(columns, (list, dict)) else columns
            rows.append([name, display(first(item, ("rows", "row_count", "n_rows"))), display(column_count), display(first(item, ("bytes", "size_bytes", "file_size")))])
            digest = first(item, ("sha256", "hash", "checksum"), None)
            if digest is not None:
                hashes.append([name, display(digest)])
            if isinstance(columns, (list, dict)):
                schemas.append([name, ", ".join(str(column) for column in columns)])
        add_table(document, ["Bảng / file", "Dòng", "Cột", "Byte"], rows)
        if hashes:
            add_table(document, ["Bảng / file", "Hash theo báo cáo"], hashes)
        if schemas:
            add_table(document, ["Bảng / file", "Tên cột theo báo cáo"], schemas)
    for key in ("metrics", "baseline", "baseline_metrics", "flight_metrics", "reconciliation"):
        if key in report:
            document.add_heading(f"Chỉ tiêu: {key}", level=2)
            add_table(document, ["Chỉ tiêu (tên gốc JSON)", "Giá trị"], list(metric_rows(report[key])))
    if isinstance(checks, list) and checks:
        document.add_heading("Danh mục kiểm tra và kết quả", level=2)
        rows = []
        for index, item in enumerate(checks):
            if isinstance(item, Mapping):
                rows.append([
                    display(first(item, ("check", "name"), f"check[{index}]")),
                    (f"{len(item['observed'])} giá trị" if isinstance(item.get("observed"), list) else display(item.get("observed", "Không có"))),
                    ("Khớp toàn bộ" if isinstance(item.get("expected"), list) and item.get("passed") else display(item.get("expected", "Không có"))),
                    ("PASS" if item.get("passed") is True else "FAIL" if item.get("passed") is False else "Không có"),
                ])
            else:
                rows.append([f"check[{index}]", display(item), "Không có", "Không có"])
        add_table(document, ["Kiểm tra", "Quan sát", "Kỳ vọng", "Kết quả"], rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--validation-report", type=Path, help="Explicit report path; missing or invalid files fail")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    report_path = args.validation_report if args.validation_report is not None else DEFAULT_REPORT
    report = None
    if args.validation_report is not None or report_path.exists():
        report = json.loads(report_path.read_text(encoding="utf-8-sig"))
        if not isinstance(report, dict):
            raise ValueError("validation_report.json phải chứa một JSON object")
    # Read every required source before writing; no partial Word when a DAX file is absent.
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    dax_sources = [(path.name, path.read_text(encoding="utf-8")) for path in DAX_FILES]
    document = Document()
    configure(document)
    render_markdown(document, readme)
    add_evidence(document, report_path, report)
    document.add_page_break()
    document.add_heading("Phụ lục B. Định nghĩa DAX để nhập từng measure", level=1)
    document.add_paragraph("Mã dưới đây được chép nguyên văn từ tệp bàn giao, chưa phải bằng chứng thực thi trong Power BI. Tạo riêng từng measure; không dán nguyên tệp như một DAX query.")
    for name, content in dax_sources:
        document.add_heading(name, level=2)
        add_code(document, content)
    document.add_paragraph(f"Word được tạo lúc {datetime.now(timezone.utc).isoformat(timespec='seconds')} (UTC); đây là thời điểm tạo tài liệu, không phải thời điểm kiểm chứng dữ liệu.")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)
    print(f"Created {args.output}")
    print(f"Evidence: {report_path if report is not None else 'not supplied; no PASS inferred'}")


if __name__ == "__main__":
    main()
